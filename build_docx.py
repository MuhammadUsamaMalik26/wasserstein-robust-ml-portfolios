import re, subprocess
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
s = open("main.tex").read()
labels = re.findall(r"\\label\{(eq:[^}]+)\}", s)
num = {l: i + 1 for i, l in enumerate(labels)}
for l, k in num.items():
    s = s.replace(f"\\label{{{l}}}", f"\\tag{{{k}}}")
    s = s.replace(f"\\eqref{{{l}}}", f"({k})")
s = s.replace(".pdf}", ".png}")
# expand table inputs with unicode minus so negatives never depend on math rendering
def inp(m):
    t = open(m.group(1)).read().replace("$-$", "−")
    return t
s = re.sub(r"\\input\{(tab_[a-z]+\.tex)\}", inp, s)
open("main_docx.tex", "w").write(s)
subprocess.run(["pandoc", "main_docx.tex", "-f", "latex", "-t", "docx", "--citeproc", "--bibliography=refs.bib",
                "--resource-path=.", "-o", "paper_raw.docx"], check=True)
d = Document("paper_raw.docx")
for sec in d.sections:
    sec.left_margin = sec.right_margin = Inches(0.9)
usable = 8.27 - 1.8
for t in d.tables:
    ncol = len(t.columns)
    first = 2.3 if ncol > 3 else 3.0
    other = (usable - first) / max(ncol - 1, 1)
    tblPr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed"); tblPr.append(lay)
    grid = t._tbl.tblGrid
    for j, gc in enumerate(grid.findall(qn("w:gridCol"))):
        gc.set(qn("w:w"), str(int((first if j == 0 else other) * 1440)))
    tw = tblPr.find(qn("w:tblW"))
    if tw is None:
        tw = OxmlElement("w:tblW"); tblPr.append(tw)
    tw.set(qn("w:type"), "dxa"); tw.set(qn("w:w"), str(int(usable * 1440)))
    for row in t.rows:
        for j, c in enumerate(row.cells):
            c.width = Inches(first if j == 0 else other)
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(8)
d.save("paper.docx")
print("ok", len(d.tables))
